#!/usr/bin/env python3
"""
cdsl_holdings_to_xls.py
=======================
Reads a CDSL Consolidated Account Statement (CAS) PDF and writes TWO Excel files
(.xls by default, .xlsx also supported):

  1. <name>_holdings.xls     every stock / ETF in the demat "HOLDING STATEMENT" tables
                             (ISIN | Security | Current Bal | Frozen | Pledge |
                              Pledge Setup | Free Bal | Market Price | Value)
  2. <name>-MF.XLS           every mutual fund in the "MUTUAL FUND UNITS HELD AS ON ..." table
                             (Scheme Name | ISIN | Folio No. | Closing Bal (Units) | NAV |
                              Cumulative Amount Invested | Valuation | Unrealised Profit/Loss |
                              Unrealised Profit/Loss %)

NPS sections are ignored.

Works with both kinds of CAS PDF:
  * PDFs that carry a text layer (plain, or produced by a scanner's OCR)
  * PDFs with NO text layer (text drawn as outlines / pure images). These are
    OCR'd with Tesseract.

Accuracy safeguards (important when OCR is involved):
  * ISINs are validated with the ISIN check digit; common OCR confusions
    (O/0, I/1, S/5, B/8 ...) are repaired only if the check digit then passes.
  * Stock rows are checked: Current Bal x Price == Value.
  * Mutual fund rows are checked three ways: Units x NAV == Valuation,
    Valuation - Invested == Profit/Loss, Profit/Loss / Invested == Profit/Loss %.
    Rows that fail are flagged in the "Check" column (yellow in the sheet).
  * Stock totals per demat account are compared with the "Portfolio Value" printed
    in the statement; the mutual fund totals are compared with the "Grand Total" row
    of the mutual fund table (see the "Reconciliation" sheet of each file).

Requirements
------------
    pip install pdfplumber xlwt            # xlwt only for .xls output
    pip install openpyxl                   # only for .xlsx output
    pip install pytesseract                # only if the PDF has no text layer
    + the Tesseract OCR program itself     # same condition
        Windows : https://github.com/UB-Mannheim/tesseract/wiki
        RHEL    : dnf install tesseract
        Ubuntu  : apt install tesseract-ocr

Usage
-----
    python cdsl_holdings_to_xls.py Kiran_2026-06.pdf
        -> Kiran_2026-06_holdings.xls  and  Kiran_2026-06-MF.XLS
    python cdsl_holdings_to_xls.py Kiran_2026-06.pdf -o my_holdings.xls
        -> my_holdings.xls  and  my_holdings-MF.xls
    python cdsl_holdings_to_xls.py cas.pdf --mf-output funds.xlsx
    python cdsl_holdings_to_xls.py cas.pdf --skip-mf          # stocks only, as before
    python cdsl_holdings_to_xls.py cas.pdf --password ABCDE1234F
    python cdsl_holdings_to_xls.py cas.pdf --pages 17-29 --force-ocr
"""

import argparse
import itertools
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor

import pdfplumber

# --------------------------------------------------------------------------
# Layout constants (x positions in points on an A4 page, 595.22 pt wide).
# Measured from the table borders of the CDSL holding statement.
# --------------------------------------------------------------------------
REF_W = 595.22
COLUMNS = [                       # (header, x_left, x_right)
    ("ISIN",                      20.6, 100.6),
    ("Security",                 100.6, 240.6),
    ("Current Bal",              240.6, 295.6),
    ("Frozen Bal",               295.6, 335.6),
    ("Pledge Bal",               335.6, 375.6),
    ("Pledge Setup Bal",         375.6, 425.6),
    ("Free Bal",                 425.6, 475.6),
    ("Market Price / Face Value", 475.6, 525.7),
    ("Value (Rs)",               525.7, 580.6),
]
C_ISIN, C_NAME, C_CUR, C_FRZ, C_PLG, C_SET, C_FREE, C_PRICE, C_VALUE = range(9)

# Mutual fund table ("MUTUAL FUND UNITS HELD AS ON ...").  Numbers are right-aligned,
# so a cell is assigned to a column by the centre of its text.
MF_COLUMNS = [                    # (header, x_left, x_right)
    ("Scheme Name",                      20.6, 111.0),
    ("ISIN",                            111.0, 172.0),
    ("Folio No.",                       172.0, 245.0),
    ("Closing Bal (Units)",             245.0, 295.0),
    ("NAV (Rs)",                        295.0, 345.0),
    ("Cumulative Amount Invested (INR)", 345.0, 412.0),
    ("Valuation (Rs)",                  412.0, 485.0),
    ("Unrealised Profit/Loss",          485.0, 545.0),
    ("Unrealised Profit/Loss (%)",      545.0, 582.0),
]
M_NAME, M_ISIN, M_FOLIO, M_UNITS, M_NAV, M_INV, M_VAL, M_PL, M_PCT = range(9)
MF_NUM_COLS = [M_UNITS, M_NAV, M_INV, M_VAL, M_PL, M_PCT]
FOLIO_RE = re.compile(r"^\d{4,}(?:/\d{1,3})?$")
MF_NUM_RE = re.compile(r"^-?\d[\d,]*(?:\.\d+)?$")

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
NUM_RE = re.compile(r"^\d[\d,]*\.\d+$")
CONFUSABLE = {  # OCR look-alikes, used ONLY when the ISIN check digit fails
    "0": "OQD", "O": "0Q", "Q": "0O", "D": "0",
    "1": "IL", "I": "1L", "L": "1I",
    "5": "S", "S": "5", "8": "B", "B": "8",
    "2": "Z", "Z": "2", "6": "G", "G": "6", "7": "T", "T": "7",
}


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
class Word:
    __slots__ = ("text", "x0", "x1", "top", "bottom", "alts")

    def __init__(self, text, x0, x1, top, bottom, alts=None):
        self.text, self.x0, self.x1, self.top, self.bottom = text, x0, x1, top, bottom
        self.alts = alts          # for OCR'd ISIN cells: all the different reads

    @property
    def xc(self):
        return (self.x0 + self.x1) / 2

    @property
    def cy(self):
        return (self.top + self.bottom) / 2


def col_of(xc):
    for i, (_, a, b) in enumerate(COLUMNS):
        if (a - 5 if i == 0 else a - 2) <= xc < b + (6 if i == len(COLUMNS) - 1 else 0):
            return i
    return None


def mf_col_of(xc):
    for i, (_, a, b) in enumerate(MF_COLUMNS):
        if a <= xc < b + (6 if i == len(MF_COLUMNS) - 1 else 0):
            return i
    return None


def group_lines(words, tol=3.0):
    """Group words into text lines by vertical centre. Returns [(cy, [words])]."""
    lines = []
    for w in sorted(words, key=lambda w: w.cy):
        if lines and abs(w.cy - lines[-1][0]) <= tol:
            lines[-1][1].append(w)
            n = len(lines[-1][1])
            lines[-1][0] += (w.cy - lines[-1][0]) / n
        else:
            lines.append([w.cy, [w]])
    for ln in lines:
        ln[1].sort(key=lambda w: w.x0)
    return lines


def isin_checksum_ok(s):
    if not ISIN_RE.match(s):
        return False
    digits = "".join(str(int(c, 36)) for c in s)
    total = 0
    for i, ch in enumerate(reversed(digits)):
        d = int(ch)
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def _valid_isin(t):
    # Indian equity ISINs (INE...) always end in 5 digits (security type, serial, check digit)
    return isin_checksum_ok(t) and (not t.startswith("INE") or t[7:].isdigit())


def _twelve_char_variants(s):
    """12-char strings obtainable from an OCR read by deleting 0..n stray characters."""
    if len(s) == 12:
        return [(s, 0)]
    if 12 < len(s) <= 15:
        out = []
        n = len(s) - 12
        for dele in itertools.combinations(range(3, len(s)), n):
            out.append(("".join(c for i, c in enumerate(s) if i not in dele), n))
        # most likely: junk at the end (e.g. the '!!' flag read as letters)
        out.sort(key=lambda x: x[0] != s[:12])
        return out
    return []


def choose_isin(reads):
    """
    Decide on one ISIN from several OCR reads of the same cell.
    A candidate must pass the ISIN check digit (and the INE structure rule).  Candidates
    needing fewer edits win; ties are broken by how many reads agree.
    Returns (isin, note); note == '' when a read passed cleanly.
    """
    norm = []
    for raw in reads:
        s = re.sub(r"[^A-Z0-9]", "", raw.upper())
        if len(s) >= 12:
            s = "IN" + s[2:]
        if s:
            norm.append(s)
    if not norm:
        return "", "ISIN not read - check manually"

    def rank_and_pick(max_subs):
        found = {}                                   # isin -> [best cost, votes]
        for s in norm:
            seen = set()
            for base, dels in _twelve_char_variants(s):
                positions = [i for i in range(2, 12) if base[i] in CONFUSABLE]
                for k in range(0, max_subs + 1):
                    for combo in itertools.combinations(positions, k):
                        for subs in itertools.product(*[CONFUSABLE[base[i]] for i in combo]):
                            t = list(base)
                            for i, ch in zip(combo, subs):
                                t[i] = ch
                            t = "".join(t)
                            if t in seen or not _valid_isin(t):
                                continue
                            seen.add(t)
                            cost = dels + k
                            rec = found.setdefault(t, [cost, 0])
                            rec[0] = min(rec[0], cost)
                            rec[1] += 1
                if seen:
                    break                           # fewest-deletion variant that works is enough
        return found

    for max_subs in (0, 1, 2, 3):
        found = rank_and_pick(max_subs)
        if found:
            ranked = sorted(found.items(), key=lambda kv: (kv[1][0], -kv[1][1]))
            best = ranked[0]
            tied = [t for t, v in ranked if v == best[1]]
            note = ""
            if best[1][0] > 0:
                note = f"ISIN repaired from OCR read '{norm[0]}'" + (" (ambiguous)" if len(tied) > 1 else "")
            return best[0], note
    return norm[0][:12], "ISIN check digit invalid - check manually"


def repair_isin(raw):
    return choose_isin([raw])


def as_number(text):
    """Parse '1,74,219.50' -> (174219.5, 2).  Returns None if not a number."""
    t = text.strip().strip("|[](){}")
    if "." in t and re.fullmatch(r"[0-9OoIl|SB,.]+", t):
        t = t.translate(str.maketrans("OoIl|SB", "0011158"))
    if NUM_RE.match(t):
        return float(t.replace(",", "")), len(t.split(".")[1])
    return None


def mf_number(text):
    """Parse a mutual fund table number: '1,23,370.81', '-863.99', '0', '24.82'.
    Returns a float or None."""
    t = text.strip().strip("|[](){}").replace("−", "-").replace("–", "-")
    if "." in t and re.fullmatch(r"-?[0-9OoIl|SB,.]+", t):
        t = t.translate(str.maketrans("OoIl|SB", "0011158"))
    if MF_NUM_RE.match(t):
        return float(t.replace(",", ""))
    return None


def clean_name(s):
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"(LIMITED|LTD)[A-Z0-9]{1,3}(?=NEW\b|UTI\b|NIPPON\b)", r"\1#", s)
    for bad, good in (("EQU ITV", "EQUITY"), ("EQUITV", "EQUITY"), ("EQU ITY", "EQUITY")):
        s = s.replace(bad, good)
    return s


# --------------------------------------------------------------------------
# Getting words (with positions) from a page: text layer or OCR
# --------------------------------------------------------------------------
def words_from_text_layer(page):
    scale = REF_W / float(page.width)
    out = []
    for w in page.extract_words(x_tolerance=2, y_tolerance=2):
        out.append(Word(w["text"], w["x0"] * scale, w["x1"] * scale, w["top"] * scale, w["bottom"] * scale))
    return out


def render(page, dpi):
    return page.to_image(resolution=dpi).original.convert("L")


def _ocr(img, x0_ref, x1_ref, config):
    """OCR one vertical strip [x0_ref, x1_ref] (in 595-pt page units) of the page image."""
    import pytesseract
    from pytesseract import Output
    k = REF_W / float(img.width)
    left = 0 if x0_ref is None else max(0, int(x0_ref / k))
    right = img.width if x1_ref is None else min(img.width, int(x1_ref / k))
    d = pytesseract.image_to_data(img.crop((left, 0, right, img.height)), lang="eng",
                                  config=config, output_type=Output.DICT)
    out = []
    for i, t in enumerate(d["text"]):
        t = t.strip()
        if not t or float(d["conf"][i]) < 0:
            continue
        x0 = (d["left"][i] + left) * k
        top = d["top"][i] * k
        out.append(Word(t, x0, x0 + d["width"][i] * k, top, top + d["height"][i] * k))
    return out


ALNUM = "ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"


def _shape_correct(boxes, scale):
    """
    Fix look-alike characters using glyph WIDTH (this font is monospaced for digits):
      * 'I' is a thin bar (~1/4 of a digit's width); '1' is as wide as other digits.
      * a letter 'O' is clearly wider (~1.3x) than the digit '0'.
    boxes: [(char, width_px)] as reported by Tesseract for one OCR'd ISIN cell.
    """
    digit_w = sorted(w for ch, w in boxes[2:] if ch in "23456789")
    ref = digit_w[len(digit_w) // 2] if digit_w else 18.0 * scale
    out = []
    for i, (ch, w) in enumerate(boxes):
        if i == 0:
            ch = "I"                       # first glyph is often merged with the cell border
        elif i == 1:
            ch = "N"
        elif ch in "1I":
            ch = "I" if w <= 0.5 * ref else "1"
        elif ch in "0OQ":
            # wide glyph = a real letter (keep Tesseract's O or Q), narrow glyph = digit zero
            ch = (ch if ch in "OQ" else "O") if w >= 1.25 * ref else "0"
        out.append(ch)
    return "".join(out)


def _isin_reads(img, cy_pt):
    """OCR one ISIN cell several ways (different magnification / segmentation)."""
    import pytesseract
    k = REF_W / float(img.width)
    box = (int(18 / k), max(0, int((cy_pt - 7.5) / k)), int(102 / k), int((cy_pt + 7.5) / k))
    cell = img.crop(box)
    reads = []
    for scale, psm in ((3, 7), (2, 7), (4, 7), (3, 8)):
        big = cell.resize((cell.width * scale, cell.height * scale))
        raw = pytesseract.image_to_boxes(
            big, lang="eng", config=f"--psm {psm} -c tessedit_char_whitelist={ALNUM}")
        boxes = []
        for line in raw.strip().splitlines():
            f = line.split()
            if len(f) >= 5:
                boxes.append((f[0], int(f[3]) - int(f[1])))
        if boxes:
            reads.append(_shape_correct(boxes, scale))
    return reads


def ocr_page(img):
    """
    Returns (row_words, page_words).
    row_words : the name column and the numeric columns are OCR'd as separate strips, and every
                ISIN cell is OCR'd on its own at several magnifications.  Far more accurate
                than OCR'ing the whole page in one go.
    page_words: a cheap full-width pass at half resolution, used only to find
                'DP Name', 'Portfolio Value' and section headings.
    """
    small = img.resize((img.width // 2, img.height // 2))
    page = _ocr(small, None, None, "--psm 6")
    # Cheap prefilter: a holdings page shows several 4-decimal prices (e.g. 515.4000).
    # Pages without them (account details, notes, ...) skip the slow strip OCR below.
    if sum(1 for w in page if re.fullmatch(r"\d[\d,]*\.\d{4}", w.text)) < 1:
        return [], page
    # NB: the space inside the quotes matters - without it Tesseract glues neighbouring numbers together
    nums = _ocr(img, 238, 585, '--psm 6 -c tessedit_char_whitelist="0123456789.,- "')
    holding_lines = []
    for cy, ws in group_lines(nums):
        parsed = [as_number(w.text) for w in ws]
        has_price = any(p and p[1] == 4 and 475 <= w.xc < 526 for p, w in zip(parsed, ws))
        has_value = any(p and p[1] == 2 and w.xc >= 525 for p, w in zip(parsed, ws))
        if has_price and has_value and sum(1 for p in parsed if p) >= 3:
            holding_lines.append(cy)
    if not holding_lines:                       # nothing to extract on this page: skip the slow part
        return [], page
    name = _ocr(img, 99, 243, "--psm 6")
    isin_words = []
    for cy in holding_lines:
        reads = _isin_reads(img, cy)
        if any(re.sub(r"[^A-Z0-9]", "", r.upper()).startswith("IN") for r in reads):
            isin_words.append(Word(reads[0], 40, 90, cy - 3, cy + 3, alts=reads))
    return isin_words + name + nums, page


# --------------------------------------------------------------------------
# Parsing one page (stocks)
# --------------------------------------------------------------------------
def find_anchors(words):
    """ISIN-looking words in the ISIN column that have numbers on the same line."""
    anchors = []
    for w in words:
        if not (20 <= w.xc <= 100.6):
            continue
        texts = w.alts or [w.text]
        if not any(11 <= len(x) <= 15 and x.startswith("IN")
                   for x in (re.sub(r"[^A-Z0-9]", "", t.upper()) for t in texts)):
            continue
        nums = [x for x in words if abs(x.cy - w.cy) <= 3.5 and x.xc >= 240 and as_number(x.text)]
        if len(nums) >= 2:
            anchors.append(w)
    anchors.sort(key=lambda w: w.cy)
    return anchors


def row_numbers(words, anchor):
    """Read the numeric cells that sit on the same line as the ISIN."""
    cells = {}
    leftovers = []
    for w in sorted(words, key=lambda w: w.x0):
        if abs(w.cy - anchor.cy) > 3.5 or w.xc < 238:
            continue
        n = as_number(w.text)
        if n is None:
            continue
        c = col_of(w.xc)
        if c is not None and c not in cells:
            cells[c] = n
        else:
            leftovers.append((w.xc, n))
    # Fall back on the number of decimals (qty=3, price=4, value=2) if a column is missing
    allnums = sorted([(w.xc, as_number(w.text)) for w in words
                      if abs(w.cy - anchor.cy) <= 3.5 and w.xc >= 238 and as_number(w.text)])
    if C_PRICE not in cells:
        p = [n for _, n in allnums if n[1] == 4]
        if p:
            cells[C_PRICE] = p[0]
    if C_VALUE not in cells:
        v = [n for _, n in allnums if n[1] == 2]
        if v:
            cells[C_VALUE] = v[-1]
    q = [n for _, n in allnums if n[1] == 3]
    if C_CUR not in cells and q:
        cells[C_CUR] = q[0]
    if C_FREE not in cells and q:
        cells[C_FREE] = q[-1]
    return {c: v[0] for c, v in cells.items()}


def resolve_numbers(cells):
    """Cross-check qty x price = value; return (cur, free, price, value, check_msg)."""
    cur, free, price, value = (cells.get(c) for c in (C_CUR, C_FREE, C_PRICE, C_VALUE))
    notes = []
    if cur is None and free is not None:
        cur = free
        notes.append("Current Bal missing, copied from Free Bal")
    if free is None and cur is not None:
        free = cur
    q = cur
    if None not in (q, price, value):
        if abs(q * price - value) <= 0.02:
            return cur, free, price, value, ("OK" if not notes else "; ".join(notes))
        if free is not None and abs(free * price - value) <= 0.02:
            return free, free, price, value, "Current Bal read wrongly, used Free Bal"
        return cur, free, price, value, "CHECK: qty x price does not equal value"
    if q is not None and price is not None and value is None:
        return cur, free, price, round(q * price, 2), "Value derived from qty x price (not read)"
    if q is not None and value is not None and price is None and q:
        return cur, free, round(value / q, 4), value, "Price derived from value / qty (not read)"
    if price is not None and value is not None and q is None and price:
        q = round(value / price, 3)
        return q, q, price, value, "Quantity derived from value / price (not read)"
    return cur, free, price, value, "CHECK: could not read all numbers"


def name_blocks(words, anchors, x_range=None, hdr_word="security", drop=("portfolio",)):
    """Assign the (wrapped, multi-line) name text to each anchor/row.
    Defaults are for the stock table; the mutual fund table passes its own
    name-column range, header word and the line prefixes to drop."""
    if x_range is None:
        x_range = (COLUMNS[C_NAME][1], COLUMNS[C_NAME][2])
    nw = [w for w in words if x_range[0] <= w.xc < x_range[1]]
    lines = [(cy, " ".join(x.text for x in ws)) for cy, ws in group_lines(nw)]
    # drop the header line(s) and the 'Portfolio Value' / 'Grand Total' bar
    hdr = [cy for cy, t in lines if t.strip().lower().startswith(hdr_word) and cy < anchors[0].cy]
    hdr_cy = max(hdr) if hdr else None
    lines = [(cy, t) for cy, t in lines
             if not t.lower().startswith(tuple(drop)) and (hdr_cy is None or cy > hdr_cy + 2)]
    gaps = sorted(b[0] - a[0] for a, b in zip(lines, lines[1:]) if b[0] - a[0] > 1)
    pitch = gaps[len(gaps) // 5] if gaps else 8.0          # typical line pitch inside one cell
    t_max = 1.3 * pitch
    cys = [a.cy for a in anchors]
    cuts = []
    for k in range(len(cys) - 1):
        seq = [cy for cy, _ in lines if cys[k] - 0.35 * pitch <= cy <= cys[k + 1] + 0.35 * pitch]
        if len(seq) >= 2:
            g, i = max((seq[i + 1] - seq[i], i) for i in range(len(seq) - 1))
            cuts.append((seq[i] + seq[i + 1]) / 2)
        else:
            cuts.append((cys[k] + cys[k + 1]) / 2)
    out = []
    n = len(anchors)
    for k in range(n):
        lo = cuts[k - 1] if k > 0 else float("-inf")
        hi = cuts[k] if k < n - 1 else float("inf")
        blk = [(cy, t) for cy, t in lines if lo < cy <= hi]
        if k == 0:           # contiguity upward from the anchor (stops at the header gap)
            keep, prev = [], cys[k]
            for cy, t in sorted([b for b in blk if b[0] <= cys[k] + 0.35 * pitch], reverse=True):
                if prev - cy <= t_max:
                    keep.append((cy, t)); prev = cy
            blk = sorted(keep + [b for b in blk if b[0] > cys[k] + 0.35 * pitch])
        if k == n - 1:       # contiguity downward from the anchor
            keep, prev = [], cys[k]
            for cy, t in sorted([b for b in blk if b[0] >= cys[k] - 0.35 * pitch]):
                if cy - prev <= t_max:
                    keep.append((cy, t)); prev = cy
            blk = sorted([b for b in blk if b[0] < cys[k] - 0.35 * pitch] + keep)
        out.append(clean_name(" ".join(t for _, t in blk)))
    return out


def page_events(words, page_no):
    """DP-name markers ('DP Name : X   BO ID : ...') and 'Portfolio Value' bars."""
    ev = []
    for cy, ws in group_lines(words):
        text = " ".join(w.text for w in ws)
        m = re.search(r"D\s*P\s*N\s*a\s*m\s*e\s*[:;]?\s*(.+?)\s+B\s*O\s*I\s*D", text, re.I)
        if m:
            ev.append((page_no, cy, "dp", re.sub(r"\s+", " ", m.group(1)).strip()))
        m = re.match(r"\s*Portfolio\s*Value\D{0,8}([\d,]+\.\d{1,2})", text, re.I)
        if m:
            ev.append((page_no, cy, "total", float(m.group(1).replace(",", ""))))
    return ev


def parse_page(words, page_no):
    anchors = find_anchors(words)
    rows = []
    if anchors:
        names = name_blocks(words, anchors)
        for a, name in zip(anchors, names):
            isin, isin_note = choose_isin(a.alts or [a.text])
            cells = row_numbers(words, a)
            cur, free, price, value, chk = resolve_numbers(cells)
            notes = [x for x in (isin_note, "" if chk == "OK" else chk) if x]
            rows.append({
                "page": page_no, "y": a.cy, "ISIN": isin, "Security": name,
                "Current Bal": cur,
                "Frozen Bal": cells.get(C_FRZ, 0.0), "Pledge Bal": cells.get(C_PLG, 0.0),
                "Pledge Setup Bal": cells.get(C_SET, 0.0),
                "Free Bal": free, "Market Price / Face Value": price, "Value (Rs)": value,
                "Check": "; ".join(notes) if notes else "OK",
            })
    return rows, page_events(words, page_no)


# --------------------------------------------------------------------------
# Whole-document driver (stocks)
# --------------------------------------------------------------------------
def parse_pages_arg(spec, n):
    if not spec:
        return list(range(1, n + 1))
    pages = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        pages += list(range(int(a), int(b or a) + 1))
    return [p for p in pages if 1 <= p <= n]


def extract(pdf_path, password=None, pages=None, dpi=300, force_ocr=False, workers=4, log=print):
    """Returns (rows, stated_totals, stop_page).  stop_page is the page where the demat
    section ended; the mutual fund table is searched from there on."""
    all_rows, events = [], []
    stop_page = None
    with pdfplumber.open(pdf_path, password=password) as pdf:
        page_list = parse_pages_arg(pages, len(pdf.pages))
        log(f"{len(pdf.pages)} pages in PDF; scanning {len(page_list)} of them")
        i = 0
        stop = False
        empty_run = 0
        while i < len(page_list) and not stop:
            chunk = page_list[i:i + workers]
            i += workers
            todo = []   # (page_no, mode, payload)
            for pno in chunk:
                pg = pdf.pages[pno - 1]
                if not force_ocr and len(pg.chars) >= 50:
                    w = words_from_text_layer(pg)
                    todo.append((pno, "text", (w, w), pg.width))
                else:
                    todo.append((pno, "ocr", render(pg, dpi), pg.width))
            ocr_items = [t for t in todo if t[1] == "ocr"]
            if ocr_items:
                try:
                    import pytesseract  # noqa: F401
                except ImportError:
                    sys.exit("This PDF page has no text layer, so OCR is needed:\n"
                             "  pip install pytesseract   (and install the Tesseract program)")
                with ThreadPoolExecutor(max_workers=workers) as ex:
                    ocr_words = list(ex.map(lambda t: ocr_page(t[2]), ocr_items))
                it = iter(ocr_words)
                todo = [(p, m, (next(it) if m == "ocr" else w), pw) for p, m, w, pw in todo]
            for pno, mode, (words, page_words), _ in todo:
                rows, _unused = parse_page(words, pno)
                ev = page_events(page_words, pno)
                log(f"  page {pno:>3} [{mode:<4}] {len(rows)} holding rows")
                all_rows += rows
                events += ev
                empty_run = 0 if rows else (empty_run + 1 if all_rows else 0)
                if empty_run >= 2:
                    stop = True          # holdings tables ended
                    if stop_page is None:
                        stop_page = pno
                joined = " ".join(w.text for w in page_words)
                if all_rows and re.search(r"MUTUAL\s+FUND\s+UNITS\s+HELD", joined, re.I):
                    stop = True          # demat section is over
                    if stop_page is None:
                        stop_page = pno
    if not all_rows:
        sys.exit("No holding rows were found. Try --force-ocr, a higher --dpi, or --pages.")

    # Attach each row to the demat account (DP) it belongs to, and collect stated totals
    merged = sorted([(r["page"], r["y"], "row", r) for r in all_rows] + events, key=lambda e: (e[0], e[1]))
    current, stated = "Unknown", {}
    for _, _, kind, payload in merged:
        if kind == "dp":
            current = payload
        elif kind == "row":
            payload["Demat Account (DP)"] = current
        elif kind == "total" and current != "Unknown":
            stated[current] = payload
    return all_rows, stated, stop_page


def reconcile(rows, stated):
    out = []
    dps = []
    for r in rows:
        if r["Demat Account (DP)"] not in dps:
            dps.append(r["Demat Account (DP)"])
    for dp in dps:
        mine = [r for r in rows if r["Demat Account (DP)"] == dp]
        s = round(sum((r["Value (Rs)"] or 0) for r in mine), 2)
        st = stated.get(dp)
        if st is None:
            status = "No 'Portfolio Value' line found to compare"
        else:
            status = "MATCHES statement" if abs(s - st) < 0.5 else f"MISMATCH (diff {round(s - st, 2)}) - check flagged rows"
        out.append((dp, len(mine), s, st, status))
    return out


# --------------------------------------------------------------------------
# Mutual fund table: parsing
# --------------------------------------------------------------------------
def ocr_page_words(img):
    """Cheap half-resolution OCR of a whole page (used to spot headings / the Grand Total row)."""
    small = img.resize((img.width // 2, img.height // 2))
    return _ocr(small, None, None, "--psm 6")


def _mf_isin_reads(img, cy_pt):
    """OCR one mutual fund ISIN cell.  The cell is two lines tall when the ISIN wraps
    ('INF209K01-' / 'WC7'), so the crop is taller and the page-segmentation mode is block mode."""
    import pytesseract
    k = REF_W / float(img.width)
    box = (int(110 / k), max(0, int((cy_pt - 11) / k)), int(173 / k), int((cy_pt + 11) / k))
    cell = img.crop(box)
    reads = []
    for scale, psm in ((3, 6), (2, 6), (4, 6), (3, 4)):
        big = cell.resize((cell.width * scale, cell.height * scale))
        raw = pytesseract.image_to_boxes(
            big, lang="eng", config=f"--psm {psm} -c tessedit_char_whitelist={ALNUM}")
        boxes = []
        for line in raw.strip().splitlines():
            f = line.split()
            if len(f) >= 5:
                boxes.append((f[0], int(f[3]) - int(f[1])))
        if boxes:
            reads.append(_shape_correct(boxes, scale))
    return reads


def ocr_mf_words(img):
    """Full-resolution OCR of a mutual fund table page, as separate strips (numbers, folio,
    scheme name) plus every ISIN cell on its own.  Returns a list of Words."""
    nums = _ocr(img, 243, 585, '--psm 6 -c tessedit_char_whitelist="0123456789.,- "')
    folio = _ocr(img, 170, 246, '--psm 6 -c tessedit_char_whitelist="0123456789/ "')
    name = _ocr(img, 19, 112, "--psm 6")
    row_lines = []
    for cy, ws in group_lines(nums):
        if sum(1 for w in ws if mf_number(w.text) is not None) >= 5:
            row_lines.append(cy)
    isin_words = []
    for cy in row_lines:
        reads = _mf_isin_reads(img, cy)
        if reads:
            isin_words.append(Word(reads[0], 130, 150, cy - 3, cy + 3, alts=reads))
    return isin_words + name + folio + nums


def find_mf_anchors(words, y_min=float("-inf")):
    """One anchor per mutual fund row: a folio number in the folio column with at least four
    numbers to its right on the same line (the numbers are vertically centred in the cell)."""
    anchors = []
    for w in sorted(words, key=lambda w: w.cy):
        if not (MF_COLUMNS[M_FOLIO][1] <= w.xc < MF_COLUMNS[M_FOLIO][2]) or w.cy <= y_min:
            continue
        if not FOLIO_RE.match(w.text.strip()):
            continue
        nums = [x for x in words if abs(x.cy - w.cy) <= 3.5 and x.xc >= MF_COLUMNS[M_UNITS][1]
                and mf_number(x.text) is not None]
        if len(nums) >= 4 and not (anchors and abs(anchors[-1].cy - w.cy) < 3):
            anchors.append(w)
    return anchors


def _merge_split_numbers(ws):
    """Scanner text layers sometimes put a space after a thousands separator ('3, 74,081.63'
    comes out as two words).  Glue such fragments back together."""
    out = []
    for w in sorted(ws, key=lambda w: w.x0):
        if out and out[-1].text.endswith(",") and w.x0 - out[-1].x1 < 8:
            p = out.pop()
            out.append(Word(p.text + w.text, p.x0, w.x1, min(p.top, w.top), max(p.bottom, w.bottom)))
        else:
            out.append(w)
    return out


def mf_row_numbers(words, anchor):
    """The six numeric cells on the anchor's line, keyed by column index."""
    line = _merge_split_numbers(
        [x for x in words if abs(x.cy - anchor.cy) <= 3.5
         and x.xc >= MF_COLUMNS[M_UNITS][1] and mf_number(x.text) is not None])
    if len(line) == len(MF_NUM_COLS):         # all six present: assign left to right
        return dict(zip(MF_NUM_COLS, (mf_number(x.text) for x in line)))
    cells = {}
    for x in line:                            # otherwise fall back on the column boundaries
        c = mf_col_of(x.xc)
        if c in MF_NUM_COLS and c not in cells:
            cells[c] = mf_number(x.text)
    return cells


def mf_folio(words, anchor):
    """Folio number: all words of the folio column on the anchor's line, glued together
    (a scanner text layer can split '477361573202/0' into '4' + '77361573202/0')."""
    ws = sorted([w for w in words if MF_COLUMNS[M_FOLIO][1] <= w.xc < MF_COLUMNS[M_FOLIO][2]
                 and abs(w.cy - anchor.cy) <= 3.5], key=lambda w: w.x0)
    return "".join(w.text.strip() for w in ws) or anchor.text.strip()


def mf_isin(words, anchor):
    """ISIN of one row.  It may wrap over two lines, so every word of the ISIN column
    within a few points of the anchor is joined in reading order."""
    cand = [w for w in words if MF_COLUMNS[M_ISIN][1] <= w.xc < MF_COLUMNS[M_ISIN][2]
            and abs(w.cy - anchor.cy) <= 9]
    for w in cand:
        if w.alts:
            return choose_isin(w.alts)
    cand.sort(key=lambda w: (round(w.cy / 3), w.x0))
    raw = "".join(w.text for w in cand)
    isin, note = choose_isin([raw])
    if "(ambiguous)" in note:
        # several different "repairs" fit the check digit, so guessing is worse than keeping
        # what was read (a printed ISIN can itself fail the check digit)
        keep = re.sub(r"[^A-Z0-9]", "", raw.upper())[:12]
        return keep, f"ISIN check digit invalid and OCR repair is ambiguous - kept as read '{keep}', verify"
    return isin, note


def resolve_mf_numbers(cells):
    """Cross-check the row's arithmetic.  Returns (units, nav, invested, value, pl, pct, notes)."""
    units, nav, inv, val, pl, pct = (cells.get(c) for c in MF_NUM_COLS)
    notes = []
    missing = [n for n, v in (("invested", inv), ("valuation", val), ("profit/loss", pl)) if v is None]
    if len(missing) == 1:                     # exactly one of the three amounts is unreadable
        if inv is None:
            inv = round(val - pl, 2)
        elif val is None:
            val = round(inv + pl, 2)
        else:
            pl = round(val - inv, 2)
        notes.append(f"{missing[0]} derived from the other two amounts (not read)")
    if None in (units, nav, inv, val, pl):
        notes.append("CHECK: could not read all numbers")
        return units, nav, inv, val, pl, pct, notes
    if abs(units * nav - val) > max(1.0, 0.0005 * abs(val)):
        notes.append("CHECK: units x NAV does not equal valuation")
    if abs(val - inv - pl) > 0.05:
        notes.append("CHECK: valuation - invested does not equal profit/loss")
    if inv:
        expected = pl / inv * 100
        if pct is None:
            pct = round(expected, 2)
            notes.append("Profit/loss % derived from profit/loss / invested (not read)")
        elif abs(expected - pct) > 0.02:
            notes.append("CHECK: profit/loss / invested does not equal profit/loss %")
    return units, nav, inv, val, pl, pct, notes


def mf_events(words):
    """(y of the 'UNITS HELD AS ON' heading or None, (invested, valuation) of the
    'Grand Total' row or None)."""
    heading_y, grand = None, None
    for cy, ws in group_lines(words):
        text = " ".join(w.text for w in ws)
        if re.search(r"UNITS\s*HELD\s*AS\s*ON", text, re.I):
            heading_y = cy
        if re.match(r"\s*Grand\s*Total\b", text, re.I):
            vals = {}
            for w in ws:
                c = mf_col_of(w.xc)
                n = mf_number(w.text)
                if c in (M_INV, M_VAL) and n is not None:
                    vals[c] = n
            if M_INV in vals and M_VAL in vals:
                grand = (vals[M_INV], vals[M_VAL])
    return heading_y, grand


def parse_mf_page(words, page_no, y_min=float("-inf")):
    anchors = find_mf_anchors(words, y_min)
    rows = []
    if anchors:
        names = name_blocks(words, anchors, x_range=(MF_COLUMNS[M_NAME][1], MF_COLUMNS[M_NAME][2]),
                            hdr_word="scheme", drop=("grand total", "page"))
        for a, name in zip(anchors, names):
            isin, isin_note = mf_isin(words, a)
            units, nav, inv, val, pl, pct, notes = resolve_mf_numbers(mf_row_numbers(words, a))
            notes = [x for x in [isin_note] + notes if x]
            rows.append({
                "page": page_no, "Scheme Name": name, "ISIN": isin,
                "Folio No.": mf_folio(words, a), "Closing Bal (Units)": units, "NAV (Rs)": nav,
                "Cumulative Amount Invested (INR)": inv, "Valuation (Rs)": val,
                "Unrealised Profit/Loss": pl, "Unrealised Profit/Loss (%)": pct,
                "Check": "; ".join(notes) if notes else "OK",
            })
    return rows


def extract_mf(pdf_path, password=None, pages=None, dpi=300, force_ocr=False, workers=4,
               start_page=None, log=print):
    """Reads the 'MUTUAL FUND UNITS HELD AS ON ...' table.  Returns (rows, printed_grand_total)
    where printed_grand_total is (invested, valuation) or None."""
    rows, grand_total = [], None
    in_table, done = False, False
    with pdfplumber.open(pdf_path, password=password) as pdf:
        page_list = [p for p in parse_pages_arg(pages, len(pdf.pages))
                     if start_page is None or p >= start_page]
        i = 0
        while i < len(page_list) and not done:
            chunk = page_list[i:i + workers]
            i += workers
            items = []          # [pno, mode, words_or_image, page_words]
            for pno in chunk:
                pg = pdf.pages[pno - 1]
                if not force_ocr and len(pg.chars) >= 50:
                    w = words_from_text_layer(pg)
                    items.append([pno, "text", w, w])
                else:
                    items.append([pno, "ocr", render(pg, dpi), None])
            ocr_items = [t for t in items if t[1] == "ocr"]
            if ocr_items:
                try:
                    import pytesseract  # noqa: F401
                except ImportError:
                    sys.exit("This PDF page has no text layer, so OCR is needed:\n"
                             "  pip install pytesseract   (and install the Tesseract program)")
                with ThreadPoolExecutor(max_workers=workers) as ex:
                    for t, pw in zip(ocr_items, ex.map(lambda t: ocr_page_words(t[2]), ocr_items)):
                        t[3] = pw
            # decide, page by page and in order, which pages belong to the table
            plan = []           # (item, heading_y, in_this_page)
            state = in_table
            for t in items:
                heading_y, grand = mf_events(t[3])
                here = state or heading_y is not None
                plan.append((t, heading_y, here))
                state = here and grand is None
            in_table = state
            heavy = [t for t, _, here in plan if here and t[1] == "ocr"]
            if heavy:
                with ThreadPoolExecutor(max_workers=workers) as ex:
                    for t, w in zip(heavy, ex.map(lambda t: ocr_mf_words(t[2]), heavy)):
                        t[2] = w
            for t, heading_y, here in plan:
                pno, mode, words, page_words = t
                if not here:
                    continue
                if mode == "ocr":
                    page_words = words
                page_rows = parse_mf_page(words, pno, heading_y if heading_y is not None else float("-inf"))
                _, grand = mf_events(page_words)
                log(f"  page {pno:>3} [{mode:<4}] {len(page_rows)} mutual fund rows")
                rows += page_rows
                if grand is not None:
                    grand_total = grand
                    done = True
    return rows, grand_total


def reconcile_mf(rows, printed):
    inv = round(sum(r["Cumulative Amount Invested (INR)"] or 0 for r in rows), 2)
    val = round(sum(r["Valuation (Rs)"] or 0 for r in rows), 2)
    out = []
    for label, s, st in (("Cumulative amount invested (INR)", inv, printed[0] if printed else None),
                         ("Valuation (Rs)", val, printed[1] if printed else None)):
        if st is None:
            status = "No 'Grand Total' row found to compare"
        else:
            status = "MATCHES statement" if abs(s - st) < 0.5 else f"MISMATCH (diff {round(s - st, 2)}) - check flagged rows"
        out.append((label, len(rows), s, st, status))
    return out


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------
HEADERS = [c[0] for c in COLUMNS] + ["Demat Account (DP)", "PDF Page", "Check"]
WIDTHS = [16, 60, 13, 11, 11, 15, 13, 24, 16, 34, 9, 60]
RECON_HEADERS = ["Demat account (DP)", "Rows read", "Sum of Value (Rs)", "Total printed in statement (Rs)", "Result"]


def row_values(r):
    return [r["ISIN"], r["Security"], r["Current Bal"], r["Frozen Bal"], r["Pledge Bal"],
            r["Pledge Setup Bal"], r["Free Bal"], r["Market Price / Face Value"], r["Value (Rs)"],
            r["Demat Account (DP)"], r["page"], r["Check"]]


def write_xls(path, rows, recon):
    try:
        import xlwt
    except ImportError:
        sys.exit("Writing .xls needs xlwt:  pip install xlwt     (or use an output name ending in .xlsx)")
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Holdings")
    head = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour gray25; align: wrap on, vert centre")
    fmts = [None, None, "#,##0.000", "#,##0.000", "#,##0.000", "#,##0.000", "#,##0.000",
            "#,##0.0000", "#,##0.00", None, "0", None]
    plain = [xlwt.easyxf(num_format_str=f) if f else xlwt.easyxf("") for f in fmts]
    warn = [xlwt.easyxf("pattern: pattern solid, fore_colour light_yellow", num_format_str=f or "General")
            for f in fmts]
    bold_total = xlwt.easyxf("font: bold on", num_format_str="#,##0.00")
    for j, h in enumerate(HEADERS):
        ws.write(0, j, h, head)
        ws.col(j).width = WIDTHS[j] * 256
    for i, r in enumerate(rows, start=1):
        style = plain if r["Check"] == "OK" else warn
        for j, v in enumerate(row_values(r)):
            ws.write(i, j, v, style[j])
    t = len(rows) + 1
    ws.write(t, 1, "TOTAL", xlwt.easyxf("font: bold on"))
    ws.write(t, 8, xlwt.Formula(f"SUM(I2:I{t})"), bold_total)
    ws.set_panes_frozen(True)
    ws.set_horz_split_pos(1)
    ws2 = wb.add_sheet("Reconciliation")
    for j, h in enumerate(RECON_HEADERS):
        ws2.write(0, j, h, head)
        ws2.col(j).width = [34, 11, 20, 30, 50][j] * 256
    money = xlwt.easyxf(num_format_str="#,##0.00")
    for i, (dp, n, s, st, status) in enumerate(recon, start=1):
        ws2.write(i, 0, dp); ws2.write(i, 1, n); ws2.write(i, 2, s, money)
        if st is not None:
            ws2.write(i, 3, st, money)
        ws2.write(i, 4, status)
    wb.save(path)


def write_xlsx(path, rows, recon):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    fmts = [None, None, "#,##0.000", "#,##0.000", "#,##0.000", "#,##0.000", "#,##0.000",
            "#,##0.0000", "#,##0.00", None, "0", None]
    for r in rows:
        ws.append(row_values(r))
        for j, f in enumerate(fmts, start=1):
            cell = ws.cell(row=ws.max_row, column=j)
            if f:
                cell.number_format = f
            if r["Check"] != "OK":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    t = ws.max_row + 1
    ws.cell(row=t, column=2, value="TOTAL").font = Font(bold=True)
    ws.cell(row=t, column=9, value=f"=SUM(I2:I{t - 1})").number_format = "#,##0.00"
    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Reconciliation")
    ws2.append(RECON_HEADERS)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for dp, n, s, st, status in recon:
        ws2.append([dp, n, s, st, status])
    for j, w in enumerate([34, 11, 20, 30, 50], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    wb.save(path)


MF_HEADERS = [c[0] for c in MF_COLUMNS] + ["PDF Page", "Check"]
MF_WIDTHS = [70, 16, 16, 14, 12, 18, 16, 16, 14, 9, 60]
MF_FMTS = [None, None, None, "#,##0.000", "#,##0.0000", "#,##0.00", "#,##0.00", "#,##0.00", "0.00", "0", None]
MF_RECON_HEADERS = ["Item", "Rows read", "Sum of rows read (Rs)", "Grand Total printed in statement (Rs)", "Result"]
MF_TOTAL_COLS = (5, 6, 7)       # invested, valuation, profit/loss (0-based) get a SUM row


def mf_row_values(r):
    return [r["Scheme Name"], r["ISIN"], r["Folio No."], r["Closing Bal (Units)"], r["NAV (Rs)"],
            r["Cumulative Amount Invested (INR)"], r["Valuation (Rs)"], r["Unrealised Profit/Loss"],
            r["Unrealised Profit/Loss (%)"], r["page"], r["Check"]]


def write_mf_xls(path, rows, recon):
    try:
        import xlwt
    except ImportError:
        sys.exit("Writing .xls needs xlwt:  pip install xlwt     (or use an output name ending in .xlsx)")
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Mutual Funds")
    head = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour gray25; align: wrap on, vert centre")
    plain = [xlwt.easyxf(num_format_str=f) if f else xlwt.easyxf("") for f in MF_FMTS]
    warn = [xlwt.easyxf("pattern: pattern solid, fore_colour light_yellow", num_format_str=f or "General")
            for f in MF_FMTS]
    bold_total = xlwt.easyxf("font: bold on", num_format_str="#,##0.00")
    for j, h in enumerate(MF_HEADERS):
        ws.write(0, j, h, head)
        ws.col(j).width = MF_WIDTHS[j] * 256
    for i, r in enumerate(rows, start=1):
        style = plain if r["Check"] == "OK" else warn
        for j, v in enumerate(mf_row_values(r)):
            ws.write(i, j, v, style[j])
    t = len(rows) + 1
    ws.write(t, 0, "TOTAL", xlwt.easyxf("font: bold on"))
    for j in MF_TOTAL_COLS:
        col = chr(ord("A") + j)
        ws.write(t, j, xlwt.Formula(f"SUM({col}2:{col}{t})"), bold_total)
    ws.set_panes_frozen(True)
    ws.set_horz_split_pos(1)
    ws2 = wb.add_sheet("Reconciliation")
    for j, h in enumerate(MF_RECON_HEADERS):
        ws2.write(0, j, h, head)
        ws2.col(j).width = [38, 11, 24, 38, 50][j] * 256
    money = xlwt.easyxf(num_format_str="#,##0.00")
    for i, (label, n, s, st, status) in enumerate(recon, start=1):
        ws2.write(i, 0, label); ws2.write(i, 1, n); ws2.write(i, 2, s, money)
        if st is not None:
            ws2.write(i, 3, st, money)
        ws2.write(i, 4, status)
    wb.save(path)


def write_mf_xlsx(path, rows, recon):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Mutual Funds"
    ws.append(MF_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for r in rows:
        ws.append(mf_row_values(r))
        for j, f in enumerate(MF_FMTS, start=1):
            cell = ws.cell(row=ws.max_row, column=j)
            if f:
                cell.number_format = f
            if r["Check"] != "OK":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    t = ws.max_row + 1
    ws.cell(row=t, column=1, value="TOTAL").font = Font(bold=True)
    for j in MF_TOTAL_COLS:
        col = get_column_letter(j + 1)
        c = ws.cell(row=t, column=j + 1, value=f"=SUM({col}2:{col}{t - 1})")
        c.number_format = "#,##0.00"
        c.font = Font(bold=True)
    for j, w in enumerate(MF_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Reconciliation")
    ws2.append(MF_RECON_HEADERS)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for row in recon:
        ws2.append(list(row))
    for j, w in enumerate([38, 11, 24, 38, 50], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    wb.save(path)


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Extract demat stock holdings and mutual fund holdings from a CDSL CAS PDF into Excel.")
    ap.add_argument("pdf", help="CDSL CAS PDF file")
    ap.add_argument("-o", "--output", help="stock holdings file (.xls or .xlsx); default: <pdf name>_holdings.xls")
    ap.add_argument("--mf-output",
                    help="mutual fund holdings file (.xls or .xlsx); default: <pdf name>-MF.XLS, "
                         "or <stock file name>-MF.<ext> when -o is given")
    ap.add_argument("--skip-mf", action="store_true", help="do not extract mutual funds (stocks only)")
    ap.add_argument("--password", help="PDF password (CDSL uses your PAN in capitals)")
    ap.add_argument("--pages", help="restrict scan, e.g. 17-23 or 18,19,20 (default: all pages)")
    ap.add_argument("--dpi", type=int, default=300, help="render resolution for OCR (default 300)")
    ap.add_argument("--force-ocr", action="store_true", help="ignore any text layer and OCR the page images")
    ap.add_argument("--workers", type=int, default=4, help="parallel OCR workers (default 4)")
    ap.add_argument("--tesseract-cmd", help="full path to tesseract executable (Windows)")
    a = ap.parse_args()

    if a.tesseract_cmd:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = a.tesseract_cmd
    out = a.output or os.path.splitext(a.pdf)[0] + "_holdings.xls"
    if a.mf_output:
        mf_out = a.mf_output
    elif a.output:
        base, ext = os.path.splitext(out)
        mf_out = base + "-MF" + ext
    else:
        mf_out = os.path.splitext(a.pdf)[0] + "-MF.XLS"

    rows, stated, stop_page = extract(a.pdf, a.password, a.pages, a.dpi, a.force_ocr, a.workers)
    recon = reconcile(rows, stated)
    (write_xlsx if out.lower().endswith(".xlsx") else write_xls)(out, rows, recon)

    print(f"\nWrote {len(rows)} stock holdings to {out}\n")
    for dp, n, s, st, status in recon:
        print(f"  {dp}: {n} rows, sum Rs {s:,.2f}"
              + (f", statement Rs {st:,.2f}" if st is not None else "") + f"  -> {status}")
    flagged = [r for r in rows if r["Check"] != "OK"]
    if flagged:
        print(f"\n{len(flagged)} row(s) flagged in the 'Check' column - please verify against the PDF:")
        for r in flagged:
            print(f"  p{r['page']} {r['ISIN']}  {r['Check']}")

    if a.skip_mf:
        return

    print("\nMutual fund table ...")
    mf_rows, printed = extract_mf(a.pdf, a.password, a.pages, a.dpi, a.force_ocr, a.workers, stop_page)
    if not mf_rows:
        print("No mutual fund rows were found (no 'MUTUAL FUND UNITS HELD AS ON' table, or it could not be "
              "read). Try --force-ocr, a higher --dpi, or --pages.")
        return
    mf_recon = reconcile_mf(mf_rows, printed)
    (write_mf_xlsx if mf_out.lower().endswith(".xlsx") else write_mf_xls)(mf_out, mf_rows, mf_recon)

    print(f"\nWrote {len(mf_rows)} mutual fund holdings to {mf_out}\n")
    for label, n, s, st, status in mf_recon:
        print(f"  {label}: sum Rs {s:,.2f}"
              + (f", statement Rs {st:,.2f}" if st is not None else "") + f"  -> {status}")
    mf_flagged = [r for r in mf_rows if r["Check"] != "OK"]
    if mf_flagged:
        print(f"\n{len(mf_flagged)} mutual fund row(s) flagged in the 'Check' column - please verify against the PDF:")
        for r in mf_flagged:
            print(f"  p{r['page']} {r['ISIN']}  {r['Check']}")


if __name__ == "__main__":
    main()
